"""Ingest KB documents into a Chroma vector store."""

from __future__ import annotations

import argparse
import os
import re
from typing import Iterable, Iterator, List, Optional, Tuple

import pandas as pd


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str, chunk_size: int, overlap: int) -> List[str]:
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive.")
    if overlap < 0:
        raise ValueError("overlap must be >= 0.")
    step = chunk_size - overlap
    if step <= 0:
        raise ValueError("overlap must be smaller than chunk_size.")
    chunks = []
    for start in range(0, len(text), step):
        chunk = text[start : start + chunk_size]
        chunk = chunk.strip()
        if chunk:
            chunks.append(chunk)
    return chunks


def read_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as handle:
        return handle.read()


def read_pdf(path: str) -> List[Tuple[str, str]]:
    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:
        raise RuntimeError("PyPDF2 is required. Install with: pip install PyPDF2") from exc

    reader = PdfReader(path)
    pages = []
    for idx, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        pages.append((text, f"page_{idx}"))
    return pages


def read_docx(path: str) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required. Install with: pip install python-docx") from exc

    doc = Document(path)
    lines = []
    for paragraph in doc.paragraphs:
        text = (paragraph.text or "").strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [((cell.text or "").strip()) for cell in row.cells]
            row_text = " | ".join([cell for cell in cells if cell])
            if row_text:
                lines.append(row_text)
    return "\n".join(lines)


def read_xlsx(path: str) -> List[Tuple[str, str]]:
    try:
        sheets = pd.read_excel(path, sheet_name=None)
    except ImportError as exc:
        raise RuntimeError("pandas is required. Install with: pip install pandas openpyxl") from exc

    outputs = []
    for name, df in sheets.items():
        df = df.fillna("")
        text = df.to_csv(index=False)
        outputs.append((text, f"sheet:{name}"))
    return outputs


def iter_documents(kb_dir: str) -> Iterator[Tuple[str, str, Optional[str]]]:
    exts = {".txt", ".md", ".pdf", ".docx", ".xlsx"}
    for root, _, files in os.walk(kb_dir):
        for fname in files:
            ext = os.path.splitext(fname)[1].lower()
            if ext in exts:
                yield os.path.join(root, fname), ext.lstrip("."), None


def build_chunks_for_file(
    path: str,
    doc_type: str,
    chunk_size: int,
    overlap: int,
) -> List[Tuple[str, Optional[str]]]:
    if doc_type in {"txt", "md"}:
        text = clean_text(read_txt(path))
        return [(chunk, None) for chunk in chunk_text(text, chunk_size, overlap)]
    if doc_type == "pdf":
        chunks = []
        for page_text, label in read_pdf(path):
            page_text = clean_text(page_text)
            for chunk in chunk_text(page_text, chunk_size, overlap):
                chunks.append((chunk, label))
        return chunks
    if doc_type == "docx":
        text = clean_text(read_docx(path))
        return [(chunk, None) for chunk in chunk_text(text, chunk_size, overlap)]
    if doc_type == "xlsx":
        chunks = []
        for sheet_text, label in read_xlsx(path):
            sheet_text = clean_text(sheet_text)
            for chunk in chunk_text(sheet_text, chunk_size, overlap):
                chunks.append((chunk, label))
        return chunks
    return []


def make_id(source_file: str, chunk_id: int, page_or_sheet: Optional[str]) -> str:
    label = page_or_sheet or "na"
    raw = f"{source_file}::{label}::{chunk_id}"
    return re.sub(r"[^a-zA-Z0-9._:-]+", "_", raw)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Ingest KB docs into Chroma")
    parser.add_argument("--kb-dir", type=str, default="kb")
    parser.add_argument("--persist-dir", type=str, default="vector_store/chroma")
    parser.add_argument(
        "--embedding-model",
        type=str,
        default="paraphrase-multilingual-MiniLM-L12-v2",
    )
    parser.add_argument("--chunk-size", type=int, default=400)
    parser.add_argument("--overlap", type=int, default=80)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    kb_dir = args.kb_dir
    if not os.path.isdir(kb_dir):
        raise SystemExit(f"kb dir not found: {kb_dir}")

    os.makedirs(args.persist_dir, exist_ok=True)

    try:
        from sentence_transformers import SentenceTransformer
    except ImportError as exc:
        raise RuntimeError(
            "sentence-transformers is required. Install with: pip install sentence-transformers"
        ) from exc

    try:
        import chromadb
    except ImportError as exc:
        raise RuntimeError("chromadb is required. Install with: pip install chromadb") from exc

    model = SentenceTransformer(args.embedding_model)
    client = chromadb.PersistentClient(path=args.persist_dir)
    collection = client.get_or_create_collection("kb_chunks")

    total_docs = 0
    total_chunks = 0

    all_texts: List[str] = []
    all_metadatas: List[dict] = []
    all_ids: List[str] = []

    for path, doc_type, _ in iter_documents(kb_dir):
        rel_path = os.path.relpath(path, kb_dir)
        chunks = build_chunks_for_file(path, doc_type, args.chunk_size, args.overlap)
        if not chunks:
            continue
        total_docs += 1
        for idx, (chunk, page_or_sheet) in enumerate(chunks):
            metadata = {
                "source_file": rel_path,
                "doc_type": doc_type,
                "chunk_id": idx,
                "page_or_sheet": page_or_sheet,
            }
            all_texts.append(chunk)
            all_metadatas.append(metadata)
            all_ids.append(make_id(rel_path, idx, page_or_sheet))
            total_chunks += 1

    if not all_texts:
        print("No chunks found to ingest.")
        return

    embeddings = model.encode(all_texts, batch_size=32, show_progress_bar=True)
    if hasattr(collection, "upsert"):
        collection.upsert(
            ids=all_ids,
            embeddings=embeddings,
            documents=all_texts,
            metadatas=all_metadatas,
        )
    else:
        collection.add(
            ids=all_ids,
            embeddings=embeddings,
            documents=all_texts,
            metadatas=all_metadatas,
        )

    print(f"Documents: {total_docs}")
    print(f"Chunks: {total_chunks}")
    print(f"Ingested into {args.persist_dir}")


if __name__ == "__main__":
    main()

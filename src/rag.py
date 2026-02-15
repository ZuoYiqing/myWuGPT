"""RAG ingestion and retrieval utilities."""

from __future__ import annotations

import argparse
import hashlib
import os
import re
from typing import Iterable, List, Optional, Tuple


DEFAULT_EMBED_MODEL = "paraphrase-multilingual-MiniLM-L12-v2"


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str, chunk_size: int, chunk_overlap: int) -> List[str]:
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive.")
    if chunk_overlap < 0:
        raise ValueError("chunk_overlap must be >= 0.")
    step = chunk_size - chunk_overlap
    if step <= 0:
        raise ValueError("chunk_overlap must be smaller than chunk_size.")
    chunks = []
    for start in range(0, len(text), step):
        chunk = text[start : start + chunk_size].strip()
        if chunk:
            chunks.append(chunk)
    return chunks


def read_txt_units(path: str) -> List[Tuple[str, str]]:
    with open(path, "r", encoding="utf-8", errors="ignore") as handle:
        text = handle.read()
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(paragraphs) <= 1:
        lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
        return [(ln, f"line_{idx}") for idx, ln in enumerate(lines, 1)]
    return [(p, f"paragraph_{idx}") for idx, p in enumerate(paragraphs, 1)]


def read_pdf_units(path: str) -> List[Tuple[str, str]]:
    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:
        raise RuntimeError("PyPDF2 is required. Install with: pip install PyPDF2") from exc

    reader = PdfReader(path)
    units = []
    for idx, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        units.append((text, f"page_{idx}"))
    return units


def read_docx_units(path: str) -> List[Tuple[str, str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required. Install with: pip install python-docx") from exc

    doc = Document(path)
    units: List[Tuple[str, str]] = []
    for idx, paragraph in enumerate(doc.paragraphs, 1):
        text = (paragraph.text or "").strip()
        if text:
            units.append((text, f"paragraph_{idx}"))
    for t_idx, table in enumerate(doc.tables, 1):
        for r_idx, row in enumerate(table.rows, 1):
            cells = [((cell.text or "").strip()) for cell in row.cells]
            row_text = " | ".join([cell for cell in cells if cell])
            if row_text:
                units.append((row_text, f"table{t_idx}_row{r_idx}"))
    return units


def read_epub_units(path: str) -> List[Tuple[str, str]]:
    try:
        from ebooklib import epub
    except ImportError:
        print("Warning: ebooklib not installed; skipping epub file.")
        return []
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        print("Warning: beautifulsoup4 not installed; skipping epub file.")
        return []

    book = epub.read_epub(path)
    units: List[Tuple[str, str]] = []
    idx = 0
    for item in book.get_items():
        if item.get_type() != epub.ITEM_DOCUMENT:
            continue
        idx += 1
        content = item.get_content().decode("utf-8", errors="ignore")
        soup = BeautifulSoup(content, "html.parser")
        text = soup.get_text(separator="\n")
        units.append((text, f"chapter_{idx}"))
    return units


def iter_kb_files(kb_dir: str) -> Iterable[Tuple[str, str]]:
    exts = {".txt", ".md", ".pdf", ".docx", ".epub"}
    for root, _, files in os.walk(kb_dir):
        for fname in files:
            ext = os.path.splitext(fname)[1].lower()
            if ext in exts:
                yield os.path.join(root, fname), ext.lstrip(".")


def load_units(path: str, doc_type: str) -> List[Tuple[str, str]]:
    if doc_type in {"txt", "md"}:
        return read_txt_units(path)
    if doc_type == "pdf":
        return read_pdf_units(path)
    if doc_type == "docx":
        return read_docx_units(path)
    if doc_type == "epub":
        return read_epub_units(path)
    return []


def make_id(source_file: str, page_or_para: str, chunk_id: int) -> str:
    safe_source = re.sub(r"[^a-zA-Z0-9._:-]+", "_", source_file)
    safe_page = re.sub(r"[^a-zA-Z0-9._:-]+", "_", page_or_para)
    digest = hashlib.sha1(source_file.encode("utf-8")).hexdigest()[:8]
    raw = f"{safe_source}::{digest}::{safe_page}::{chunk_id}"
    return raw


def split_long_paragraph(text: str, chunk_size: int, chunk_overlap: int) -> List[str]:
    return chunk_text(text, chunk_size, chunk_overlap)


def overlap_paragraphs(
    paragraphs: List[str], labels: List[str], chunk_overlap: int
) -> Tuple[List[str], List[str]]:
    if chunk_overlap <= 0:
        return [], []
    total = 0
    out_paras: List[str] = []
    out_labels: List[str] = []
    for text, label in reversed(list(zip(paragraphs, labels))):
        out_paras.insert(0, text)
        out_labels.insert(0, label)
        total += len(text)
        if total >= chunk_overlap:
            break
    return out_paras, out_labels


def build_chunks_from_units(
    units: List[Tuple[str, str]], chunk_size: int, chunk_overlap: int
) -> List[Tuple[str, str]]:
    # Merge short "title" paragraphs with the following body paragraph.
    merged_units: List[Tuple[str, str]] = []
    i = 0
    while i < len(units):
        text, label = units[i]
        if i + 1 < len(units):
            next_text, next_label = units[i + 1]
            if len(text.strip()) <= 60 and len(next_text.strip()) > 60:
                merged_units.append((f"{text}\n{next_text}", label))
                i += 2
                continue
        merged_units.append((text, label))
        i += 1

    chunks: List[Tuple[str, str]] = []
    current_paras: List[str] = []
    current_labels: List[str] = []
    current_len = 0

    def flush():
        nonlocal current_paras, current_labels, current_len
        if not current_paras:
            return
        chunk_text_value = "\n".join(current_paras).strip()
        if chunk_text_value:
            chunks.append((chunk_text_value, current_labels[0]))
        current_paras = []
        current_labels = []
        current_len = 0

    for para_text, label in merged_units:
        para_text = clean_text(para_text)
        if not para_text:
            continue
        if len(para_text) > chunk_size:
            flush()
            for seg in split_long_paragraph(para_text, chunk_size, chunk_overlap):
                if seg:
                    chunks.append((seg, label))
            continue

        if current_len and current_len + len(para_text) + 1 > chunk_size:
            overlap_paras, overlap_labels = overlap_paragraphs(
                current_paras, current_labels, chunk_overlap
            )
            flush()
            if overlap_paras:
                current_paras = overlap_paras
                current_labels = overlap_labels
                current_len = sum(len(p) for p in current_paras) + max(
                    0, len(current_paras) - 1
                )

        current_paras.append(para_text)
        current_labels.append(label)
        current_len += len(para_text) + (1 if current_len else 0)

    flush()
    return chunks


def ingest(
    kb_dir: str,
    persist_dir: str,
    chunk_size: int,
    chunk_overlap: int,
    embedding_model: str = DEFAULT_EMBED_MODEL,
) -> None:
    if not os.path.isdir(kb_dir):
        raise ValueError(f"kb_dir not found: {kb_dir}")
    os.makedirs(persist_dir, exist_ok=True)

    try:
        import chromadb
    except ImportError as exc:
        raise RuntimeError("chromadb is required. Install with: pip install chromadb") from exc

    try:
        from sentence_transformers import SentenceTransformer
    except ImportError as exc:
        raise RuntimeError(
            "sentence-transformers is required. Install with: pip install sentence-transformers"
        ) from exc

    model = SentenceTransformer(embedding_model)
    client = chromadb.PersistentClient(path=persist_dir)
    collection = client.get_or_create_collection("rag_chunks")

    all_texts: List[str] = []
    all_metadatas: List[dict] = []
    all_ids: List[str] = []
    doc_count = 0

    for path, doc_type in iter_kb_files(kb_dir):
        rel_path = os.path.relpath(path, kb_dir)
        units = load_units(path, doc_type)
        if not units:
            continue
        doc_count += 1
        if doc_type in {"pdf", "epub"}:
            for unit_text, page_or_para in units:
                cleaned = clean_text(unit_text)
                for chunk_id, chunk in enumerate(
                    chunk_text(cleaned, chunk_size, chunk_overlap)
                ):
                    metadata = {
                        "source_file": rel_path,
                        "doc_type": doc_type,
                        "page_or_para": page_or_para,
                        "chunk_id": chunk_id,
                    }
                    if page_or_para.startswith("page_"):
                        metadata["page"] = page_or_para
                    else:
                        metadata["paragraph"] = page_or_para
                    all_texts.append(chunk)
                    all_metadatas.append(metadata)
                    all_ids.append(make_id(rel_path, page_or_para, chunk_id))
        else:
            chunks = build_chunks_from_units(units, chunk_size, chunk_overlap)
            for chunk_id, (chunk, page_or_para) in enumerate(chunks):
                metadata = {
                    "source_file": rel_path,
                    "doc_type": doc_type,
                    "page_or_para": page_or_para,
                    "chunk_id": chunk_id,
                }
                if page_or_para.startswith("page_"):
                    metadata["page"] = page_or_para
                else:
                    metadata["paragraph"] = page_or_para
                all_texts.append(chunk)
                all_metadatas.append(metadata)
                all_ids.append(make_id(rel_path, page_or_para, chunk_id))

    if not all_texts:
        print("No chunks found to ingest.")
        return

    embeddings = model.encode(all_texts, batch_size=32, show_progress_bar=True)
    if hasattr(collection, "upsert"):
        collection.upsert(
            ids=all_ids,
            embeddings=embeddings,
            documents=all_texts,
            metadatas=all_metadatas,
        )
    else:
        collection.add(
            ids=all_ids,
            embeddings=embeddings,
            documents=all_texts,
            metadatas=all_metadatas,
        )

    print(f"Documents: {doc_count}")
    print(f"Chunks: {len(all_texts)}")
    if all_texts:
        print("Sample chunk:")
        print(all_texts[0][:500])
    total_size = 0
    for root, _, files in os.walk(persist_dir):
        for fname in files:
            total_size += os.path.getsize(os.path.join(root, fname))
    print(f"Persist dir size: {total_size} bytes")
    print(f"Ingested into {persist_dir}")


def retrieve(query: str, persist_dir: str, top_k: int) -> List[dict]:
    if not os.path.isdir(persist_dir):
        raise ValueError(f"persist_dir not found: {persist_dir}")

    try:
        import chromadb
    except ImportError as exc:
        raise RuntimeError("chromadb is required. Install with: pip install chromadb") from exc

    try:
        from sentence_transformers import SentenceTransformer
    except ImportError as exc:
        raise RuntimeError(
            "sentence-transformers is required. Install with: pip install sentence-transformers"
        ) from exc

    model = SentenceTransformer(DEFAULT_EMBED_MODEL)
    client = chromadb.PersistentClient(path=persist_dir)
    collection = client.get_or_create_collection("rag_chunks")

    query_embedding = model.encode([query])[0]
    result = collection.query(
        query_embeddings=[query_embedding],
        n_results=max(1, top_k),
        include=["documents", "metadatas", "distances"],
    )

    items = []
    docs = result.get("documents", [[]])[0]
    metas = result.get("metadatas", [[]])[0]
    for text, meta in zip(docs, metas):
        items.append(
            {
                "text": text,
                "metadata": meta or {},
            }
        )
    return items


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="RAG ingestion and retrieval")
    parser.add_argument("--ingest", action="store_true", help="Ingest KB into Chroma")
    parser.add_argument("--kb-dir", type=str, default="knowledge_base")
    parser.add_argument("--persist-dir", type=str, default=".chroma")
    parser.add_argument("--chunk-size", type=int, default=400)
    parser.add_argument("--chunk-overlap", type=int, default=80)
    parser.add_argument("--embedding-model", type=str, default=DEFAULT_EMBED_MODEL)
    parser.add_argument("--query", type=str, default=None)
    parser.add_argument("--top-k", type=int, default=4)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.ingest:
        ingest(
            args.kb_dir,
            args.persist_dir,
            args.chunk_size,
            args.chunk_overlap,
            args.embedding_model,
        )
        return
    if args.query:
        items = retrieve(args.query, args.persist_dir, args.top_k)
        for idx, item in enumerate(items, 1):
            meta = item.get("metadata", {})
            source = meta.get("source_file", "unknown")
            page = meta.get("page_or_para", "")
            label = f"[{idx}] {source}"
            if page:
                label += f" ({page})"
            print(label)
            print(item.get("text", ""))
            print("-" * 60)
        return
    print("Nothing to do. Use --ingest or --query.")


if __name__ == "__main__":
    main()
